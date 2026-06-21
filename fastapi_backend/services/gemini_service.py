import os
import json
import re
import asyncio
import pdfplumber
import pymupdf4llm
from dotenv import load_dotenv

load_dotenv()

# ─────────────────────────────────────────────
# Cấu hình Gemini API (dùng google-genai mới)
# ─────────────────────────────────────────────
try:
    from google import genai as _genai_new
    _client = _genai_new.Client(api_key=os.getenv("GEMINI_API_KEY"))
    _USE_NEW_SDK = True
except ImportError:
    # Fallback về SDK cũ nếu chưa cài google-genai
    import google.generativeai as genai
    genai.configure(api_key=os.getenv("GEMINI_API_KEY"))
    _USE_NEW_SDK = False


def _generate_content(model_name: str, prompt: str) -> str:
    """Gọi Gemini API – tự động chọn SDK đang có."""
    if _USE_NEW_SDK:
        response = _client.models.generate_content(model=model_name, contents=prompt)
        return response.text
    else:
        model = genai.GenerativeModel(model_name)
        return model.generate_content(prompt).text

def parse_pdf_to_text(file_path: str) -> str:
    """Đọc file PDF và trả về toàn bộ text (Legacy)."""
    text = ""
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        print(f"Error reading PDF: {e}")
    return text

def extract_schedules_from_text(text: str, departments: list) -> list:
    """
    Gửi nội dung chữ lên Gemini để bóc tách thành danh sách JSON (Legacy/Đồng bộ).
    """
    dept_info = json.dumps(departments, ensure_ascii=False)
    prompt = f"""
Bạn là một trợ lý AI phân tích lịch công tác. Nhiệm vụ của bạn là đọc nội dung lịch (dạng text được trích xuất từ PDF) và trích xuất thành một mảng (array) các object JSON.

DANH SÁCH PHÒNG BAN TRONG HỆ THỐNG:
{dept_info}

Quy tắc trích xuất cho mỗi lịch (trả về đúng định dạng JSON này, không có markdown formatting bên ngoài, CHỈ TRẢ VỀ JSON ARRAY TRỰC TIẾP `[...]`):
[
  {{
    "title": "Nội dung công việc (chuỗi)",
    "teacher": "Người chủ trì (tên người, ví dụ: Đ/c Vũ (PGĐ))",
    "room": "Địa điểm (nếu có, không có để chuỗi rỗng)",
    "scheduleDate": "Ngày diễn ra (định dạng YYYY-MM-DD). Hãy tự suy luận ngày dựa vào tiêu đề Lịch tuần hoặc các dấu hiệu ngày tháng.",
    "startTime": "Giờ bắt đầu (định dạng HH:MM, nếu không có để 08:00)",
    "endTime": "Giờ kết thúc (định dạng HH:MM, nếu không có để 11:30)",
    "note": "Thành phần dự hoặc ghi chú (chuỗi)",
    "unit": "Học viện ANND",
    "departmentId": "Mã UUID của phòng ban liên quan nhất (Dựa vào chữ viết tắt trong Nội dung hoặc Thành phần dự. Ví dụ: QLĐT, HC, NV1... Đối chiếu với DANH SÁCH PHÒNG BAN ở trên để lấy ra id chính xác. Nếu không xác định được, hãy lấy id của phòng Quản lý đào tạo (QLĐT) hoặc Hành chính (HC) hoặc để chuỗi rỗng).",
    "category": "ToanTruong",
    "participantUserIds": []
  }}
]

NỘI DUNG LỊCH CÔNG TÁC:
{text}
"""
    try:
        result_text = _generate_content('gemini-2.5-flash', prompt).strip()
        if result_text.startswith("```json"):
            result_text = result_text[7:]
        if result_text.startswith("```"):
            result_text = result_text[3:]
        if result_text.endswith("```"):
            result_text = result_text[:-3]
        result_text = result_text.strip()
        parsed_json = json.loads(result_text)
        if isinstance(parsed_json, list):
            return parsed_json
        return []
    except Exception as e:
        print(f"Gemini API Error: {e}")
        return []

# ==================== PHẦN MỚI: TỐI ƯU SONG SONG MARKDOWN ====================

def find_matching_department_id(dept_keyword: str, departments: list) -> str:
    """Ánh xạ tên phòng ban từ AI trích xuất sang mã UUID tương ứng trong cơ sở dữ liệu."""
    if not dept_keyword or not isinstance(dept_keyword, str):
        return ""
    clean_kw = dept_keyword.strip().lower()
    
    # 1. Khớp trực tiếp (ví dụ: "QLĐT" nằm trong "Quản lý đào tạo (QLĐT)")
    for dept in departments:
        dept_name = dept["name"].lower()
        if clean_kw in dept_name:
            return dept["id"]
            
    # 2. Khớp ký tự viết tắt (ví dụ: "hc" -> "hành chính")
    for dept in departments:
        dept_name = dept["name"].lower()
        words = dept_name.split()
        initials = "".join([w[0] for w in words if w])
        if clean_kw == initials:
            return dept["id"]
            
    # 3. Khớp từ khóa cụ thể trong tên
    for dept in departments:
        dept_name = dept["name"].lower()
        if clean_kw in dept_name or dept_name in clean_kw:
            return dept["id"]
            
    # 4. Fallback về Quản lý đào tạo hoặc Hành chính
    for dept in departments:
        name_lower = dept["name"].lower()
        if "đào tạo" in name_lower or "qldt" in name_lower:
            return dept["id"]
    for dept in departments:
        name_lower = dept["name"].lower()
        if "hành chính" in name_lower or "hc" in name_lower:
            return dept["id"]
            
    return departments[0]["id"] if departments else ""

def parse_pdf_to_markdown(file_path: str) -> str:
    """Chuyển đổi file PDF sang Markdown giữ nguyên định dạng bảng biểu."""
    try:
        # Sử dụng pymupdf4llm để bóc tách Markdown Table cực tốt
        return pymupdf4llm.to_markdown(file_path)
    except Exception as e:
        print(f"Error parsing PDF to Markdown: {e}")
        # Fallback về pdfplumber nếu pymupdf4llm gặp lỗi
        return parse_pdf_to_text(file_path)

def split_markdown_by_days(md_text: str) -> dict:
    """
    Tách nội dung Markdown lịch theo các Thứ trong tuần.
    Trả về dict dạng: {"Thứ Hai": "nội dung...", "Thứ Ba": "nội dung..."}
    """
    # Các mốc thứ thông dụng để cắt nhỏ
    days = ["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "Chủ Nhật", 
            "Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7", "Chủ nhật"]
    
    # Tìm kiếm các tiêu đề ngày trong Markdown sử dụng Regex
    pattern = r"(?i)(" + "|".join([re.escape(day) for day in days]) + r")"
    matches = list(re.finditer(pattern, md_text))
    
    if not matches:
        # Không tìm thấy mốc Thứ nào, giữ nguyên cả cục
        return {"Toàn bộ lịch": md_text}
    
    chunks = {}
    for i in range(len(matches)):
        start = matches[i].start()
        # Vị trí kết thúc là điểm bắt đầu của ngày tiếp theo, hoặc hết chuỗi
        end = matches[i+1].start() if i + 1 < len(matches) else len(md_text)
        day_name = matches[i].group(0)
        chunks[day_name] = md_text[start:end]
        
    return chunks

def split_markdown_into_groups(md_text: str) -> dict:
    """
    Tách lịch và gộp thành 3 nhóm để tránh giới hạn rate limit 5 RPM của Gemini API free tier.
    - Nhóm 1: Thứ 2 và Thứ 3.
    - Nhóm 2: Thứ 4 và Thứ 5.
    - Nhóm 3: Thứ 6, Thứ 7, Chủ Nhật.
    """
    chunks = split_markdown_by_days(md_text)
    
    group_mapping = {
        "thứ hai": "Nhóm 1 (Thứ 2 - Thứ 3)",
        "thứ 2": "Nhóm 1 (Thứ 2 - Thứ 3)",
        "thứ ba": "Nhóm 1 (Thứ 2 - Thứ 3)",
        "thứ 3": "Nhóm 1 (Thứ 2 - Thứ 3)",
        
        "thứ tư": "Nhóm 2 (Thứ 4 - Thứ 5)",
        "thứ 4": "Nhóm 2 (Thứ 4 - Thứ 5)",
        "thứ năm": "Nhóm 2 (Thứ 4 - Thứ 5)",
        "thứ 5": "Nhóm 2 (Thứ 4 - Thứ 5)",
        
        "thứ sáu": "Nhóm 3 (Thứ 6 - Chủ Nhật)",
        "thứ 6": "Nhóm 3 (Thứ 6 - Chủ Nhật)",
        "thứ bảy": "Nhóm 3 (Thứ 6 - Chủ Nhật)",
        "thứ 7": "Nhóm 3 (Thứ 6 - Chủ Nhật)",
        "chủ nhật": "Nhóm 3 (Thứ 6 - Chủ Nhật)",
        "chủ nhật": "Nhóm 3 (Thứ 6 - Chủ Nhật)"
    }
    
    groups = {
        "Nhóm 1 (Thứ 2 - Thứ 3)": "",
        "Nhóm 2 (Thứ 4 - Thứ 5)": "",
        "Nhóm 3 (Thứ 6 - Chủ Nhật)": ""
    }
    
    for day_name, content in chunks.items():
        matched_group = None
        for key, group_name in group_mapping.items():
            if key in day_name.lower():
                matched_group = group_name
                break
        if matched_group:
            groups[matched_group] += content + "\n"
        else:
            # Fallback nếu tên ngày lạ, gộp vào Nhóm 1
            groups["Nhóm 1 (Thứ 2 - Thứ 3)"] += content + "\n"
            
    # Lọc bỏ các nhóm rỗng
    return {k: v for k, v in groups.items() if v.strip()}

async def extract_single_chunk(group_name: str, chunk_text: str, departments: list) -> list:
    """Gọi Gemini xử lý song song trong Thread Pool để tránh blocking luồng chính."""
    
    prompt = f"""
Bạn là một trợ lý AI phân tích lịch công tác. Nhiệm vụ của bạn là đọc nội dung lịch dưới đây (dạng Markdown) và trích xuất thành danh sách các object JSON rút gọn.

Quy tắc trích xuất (CHỈ trả về JSON array trực tiếp `[...]`, không bọc trong markdown code block, không giải thích thêm):
[
  {{
    "title": "Nội dung công việc (chuỗi)",
    "teacher": "Người chủ trì (tên người, ví dụ: Đ/c Vũ (PGĐ))",
    "room": "Địa điểm (nếu có, không có để chuỗi rỗng)",
    "scheduleDate": "Ngày diễn ra (định dạng YYYY-MM-DD). Hãy suy luận ngày dựa vào mốc thời gian trong văn bản của từng công việc.",
    "startTime": "Giờ bắt đầu (định dạng HH:MM, mặc định 08:00)",
    "endTime": "Giờ kết thúc (định dạng HH:MM, mặc định 11:30)",
    "note": "Thành phần dự hoặc ghi chú (chuỗi)",
    "department": "Tên viết tắt hoặc từ khóa của phòng ban liên quan nhất (ví dụ: QLĐT, HC, NV1... để trống nếu không rõ)"
  }}
]

NỘI DUNG LỊCH CÔNG TÁC CỦA {group_name}:
{chunk_text}
"""
    try:
        # Chạy đồng bộ SDK Gemini trong Thread Pool bằng asyncio
        loop = asyncio.get_event_loop()

        # Sử dụng gemini-2.5-flash-lite để có quota lớn và tốc độ siêu tốc
        response = await loop.run_in_executor(
            None,
            lambda: _generate_content('gemini-2.5-flash-lite', prompt)
        )
        
        result_text = response.strip() if isinstance(response, str) else response
        
        # Dọn sạch các ký tự bao bọc của code block json nếu có
        if result_text.startswith("```json"):
            result_text = result_text[7:]
        if result_text.startswith("```"):
            result_text = result_text[3:]
        if result_text.endswith("```"):
            result_text = result_text[:-3]
            
        result_text = result_text.strip()
        parsed_json = json.loads(result_text)
        
        if isinstance(parsed_json, list):
            # Hậu xử lý trên Python để điền các trường mặc định và map department UUID
            for item in parsed_json:
                item["unit"] = "Học viện ANND"
                item["category"] = "ToanTruong"
                item["participantUserIds"] = []
                
                # Ánh xạ tên phòng ban sang UUID
                dept_name = item.pop("department", "")
                item["departmentId"] = find_matching_department_id(dept_name, departments)
                
            return parsed_json
        return []
    except Exception as e:
        print(f"Error extracting schedules for {group_name}: {e}")
        return []

def parse_docx_to_markdown(file_path: str) -> str:
    """Chuyển đổi file Word (.docx) sang Markdown bao gồm cả bảng biểu."""
    try:
        import docx
        doc = docx.Document(file_path)
        content = []
        
        # Duyệt tuần tự các phần tử (đoạn văn & bảng biểu) trong body của tài liệu
        for block in doc.element.body:
            name = block.tag.split('}')[-1]
            if name == 'p':
                p = docx.text.paragraph.Paragraph(block, doc)
                if p.text.strip():
                    content.append(p.text.strip())
            elif name == 'tbl':
                table = docx.table.Table(block, doc)
                table_md = []
                for r_idx, row in enumerate(table.rows):
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    table_md.append(f"| {' | '.join(row_cells)} |")
                    if r_idx == 0:
                        separators = ["---"] * len(row_cells)
                        table_md.append(f"| {' | '.join(separators)} |")
                content.append("\n".join(table_md) + "\n")
        return "\n".join(content)
    except Exception as e:
        print(f"Error parsing Word (.docx) to Markdown: {e}")
        return ""

def parse_xlsx_to_markdown(file_path: str) -> str:
    """Chuyển đổi file Excel (.xlsx) sang Markdown Table."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheet = wb.active
        content = []
        
        table_md = []
        for r_idx, row in enumerate(sheet.iter_rows(values_only=True)):
            # Chuyển các cell None thành chuỗi rỗng và lọc bỏ ký tự xuống dòng
            row_cells = [str(cell).strip().replace("\n", " ") if cell is not None else "" for cell in row]
            
            # Bỏ qua các dòng trống hoàn toàn
            if not any(row_cells):
                continue
                
            table_md.append(f"| {' | '.join(row_cells)} |")
            if r_idx == 0:
                separators = ["---"] * len(row_cells)
                table_md.append(f"| {' | '.join(separators)} |")
                
        content.append("\n".join(table_md))
        wb.close()
        return "\n".join(content)
    except Exception as e:
        print(f"Error parsing Excel (.xlsx) to Markdown: {e}")
        return ""

async def extract_schedules_from_file_async(file_path: str, file_ext: str, departments: list) -> list:
    """Hàm chính điều phối: Trích xuất nội dung tùy định dạng -> Phân tách theo Thứ -> Nhóm 3 luồng song song -> Gộp kết quả."""
    if file_ext == '.pdf':
        md_text = parse_pdf_to_markdown(file_path)
    elif file_ext == '.docx':
        md_text = parse_docx_to_markdown(file_path)
    elif file_ext == '.xlsx':
        md_text = parse_xlsx_to_markdown(file_path)
    else:
        print(f"Unsupported file extension: {file_ext}")
        return []
        
    if not md_text or not md_text.strip():
        print("Empty text extracted from file.")
        return []
        
    # Phân tách và gộp nội dung lịch thành 3 nhóm (Tránh rate limit 5 RPM)
    groups = split_markdown_into_groups(md_text)
    
    # Tạo danh sách task chạy song song bằng asyncio (Tối đa 3 tasks concurrent)
    tasks = []
    for group_name, chunk_content in groups.items():
        tasks.append(extract_single_chunk(group_name, chunk_content, departments))
        
    # Thực thi song song và chờ tất cả hoàn thành
    results = await asyncio.gather(*tasks)
    
    # Gộp danh sách lịch từ các nhóm ngày
    all_schedules = []
    for group_schedules in results:
        all_schedules.extend(group_schedules)
        
    return all_schedules

